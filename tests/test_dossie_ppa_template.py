"""Dossiê Desalinhado: o gerador tem de acompanhar o modelo oficial.

O modelo trocou (`Parecer_Invista_PPA_v2_Atualizada_3.0` → `Parecer_Desalinhado_PPA`) e
mexeu justamente no que o código preenchia: renomeou "Dados do Processo" para "Resumo do
Processo", tirou o lastro de lá e o promoveu a quadro repetível, dividiu os índices do
contrato em adimplemento/inadimplemento, acrescentou REDES SOCIAIS no meio da seção 0 e
passou a chamar VP de VF.

Estes testes existem porque nada disso quebra com erro: um rótulo que deixou de existir
faz o quadro sair vazio, em silêncio, num documento que continua abrindo.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402

import dossie_ppa  # noqa: E402


DADOS = {
    "nome_caso": "Banco do Brasil x Agropecuária Teste",
    "exequentes": "Banco do Brasil S.A.",
    "executados": "Agropecuária Teste Ltda",
    "sat_total": "R$ 4.200.000,00",
    "total_atingivel_vm": "R$ 1.980.000,00",
    "total_atingivel_vp": "R$ 1.400.000,00",
    "consideracoes_gerais": "Execução com penhora deferida.",
    "redes_sociais": [
        {"plataforma": "Instagram", "link": "@agroteste", "observacoes": "Perfil da executada (fls. 320)"},
        {"plataforma": "Kwai", "link": "kwai.com/@joao", "observacoes": "Perfil do sócio (fls. 322)"},
    ],
    "visao_consolidada_ativos": [
        {"tese": "Penhora Direta", "vm": "R$ 1.500.000,00", "vp": "R$ 1.000.000,00",
         "onus": "R$ 300.000,00"},
    ],
    "ativos": [
        {"tese": "Penhora Direta", "matricula": "Matrícula nº 30.174", "comarca": "Curitiba/PR",
         "proprietario_atual": "Agropecuária Teste Ltda", "descricao": "Lotes 08 a 13",
         "vm": "R$ 1.500.000,00", "vp": "R$ 1.000.000,00"},
    ],
    "teses_recuperacao": {
        "penhora_direta": {"analise": "Penhora deferida sobre a sede."},
        "idpj": {"resumo": "Confusão patrimonial com a holding.",
                 "empresa_alvo": {"razao_social": "Holding Teste Ltda", "cnpj": "22.333.444/0001-55"}},
        "fraude_execucao": {"resumo": "Doação ao filho após a citação."},
    },
    "creditos": [{
        "id": "Crédito Banco do Brasil",
        "numero_processo": "0001234-56.2019.8.16.0014",
        "vara_comarca": "1ª Vara Cível — Curitiba/PR",
        "prescricao": "Intercorrente em curso desde 12/2021 (Mov. 30)",
        "sucumbencia": "Executada condenada em 10% (Mov. 51)",
        "riscos_juridicos": "Citação por edital questionável (fls. 40).",
        "status_processo": "Em fase de penhora (Mov. 51)",
        "lastros": [
            {"lastro": "CCB nº 40/00123-4", "data_emissao": "10/01/2018",
             "data_vencimento": "10/01/2020", "partes": "Agropecuária Teste Ltda (emitente)",
             "destinacao": "Custeio de safra", "assinaturas": "João da Silva (fls. 20)",
             "garantia": "Hipoteca cedular sobre a matrícula 30.174"},
            {"lastro": "CPR nº 88/2019", "data_emissao": "05/05/2019",
             "destinacao": "Capital de giro", "garantia": "Penhor de safra"},
        ],
        "adimp_cm": "IPCA", "adimp_jr": "1% ao mês", "adimp_cap": "mensal",
        "ind_cm": "INPC", "ind_jr": "1,5% ao mês", "ind_jm": "1% ao mês",
        "ind_multa": "2%", "ind_cap": "mensal", "ind_comissao": "Vedada cumulação",
        "plan_cm": "INPC", "plan_comissao": "Não aplicada",
    }],
}


def _tabela(doc, prefixo: str):
    tabela = dossie_ppa._tabela_por_titulo(doc, prefixo)
    assert tabela is not None, f"quadro não localizado no modelo: {prefixo}"
    return tabela


def _valores(tabela) -> dict:
    return {linha.cells[0].text.strip(): linha.cells[1].text.strip() for linha in tabela.rows}


def _texto_inteiro(doc) -> str:
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            partes.extend(c.text for c in linha.cells)
    return "\n".join(partes)


@pytest.fixture(scope="module")
def documento():
    return Document(dossie_ppa._build_doc(DADOS))


def testar_modelo_desalinhado_vai_junto_no_deploy():
    assert dossie_ppa._TEMPLATE_PATH.exists(), "o modelo precisa ir para o Space"
    assert dossie_ppa._TEMPLATE_PATH.name == "Parecer_Desalinhado_PPA.docx"


def testar_resumo_do_processo_substituiu_dados_do_processo(documento):
    """O quadro mudou de nome e trocou o bloco do lastro por prescrição/sucumbência/riscos."""
    valores = _valores(_tabela(documento, "Resumo do Processo"))
    assert valores["Vara / Comarca"] == "1ª Vara Cível — Curitiba/PR"
    assert valores["Prescrição?"].startswith("Intercorrente em curso")
    assert valores["Sucumbência?"].startswith("Executada condenada")
    assert valores["Riscos jurídicos gerais"].startswith("Citação por edital")
    assert "Lastro / Instrumento" not in valores, "o lastro saiu deste quadro"


def testar_um_quadro_por_lastro(documento):
    """Uma execução pode cobrar mais de um título — o quadro é repetível."""
    titulos = [p.text.strip() for p in documento.paragraphs if p.text.strip().startswith("Lastro nº")]
    assert titulos == ["Lastro nº 1", "Lastro nº 2"]
    primeiro = _valores(_tabela(documento, "Lastro nº 1"))
    assert primeiro["Lastro / Instrumento"] == "CCB nº 40/00123-4"
    assert primeiro["Partes"] == "Agropecuária Teste Ltda (emitente)"
    assert primeiro["Destinação do Lastro"] == "Custeio de safra"
    # Campo vazio no 2º lastro não herda o valor do 1º.
    segundo = _valores(_tabela(documento, "Lastro nº 2"))
    assert segundo["Lastro / Instrumento"] == "CPR nº 88/2019"
    assert segundo["Data do Vencimento"] == ""


def testar_lastro_no_formato_antigo_continua_abrindo():
    """Extração e dossiês salvos antes do modelo novo traziam o lastro solto."""
    antigo = {"lastro": "CCB nº 1", "data_emissao": "01/01/2020", "garantia": "Hipoteca"}
    assert dossie_ppa.lastros_do_credito(antigo) == [{
        "lastro": "CCB nº 1", "data_emissao": "01/01/2020", "data_vencimento": "",
        "assinaturas": "", "garantia": "Hipoteca",
    }]
    assert dossie_ppa.lastros_do_credito({}) == []
    assert dossie_ppa.lastros_do_credito({"lastros": [{"lastro": "CPR"}]}) == [{"lastro": "CPR"}]


def testar_indices_separam_adimplemento_de_inadimplemento(documento):
    """São regimes distintos: o que o contrato cobra em dia não é o que cobra em atraso."""
    adimplemento = _valores(_tabela(documento, "Índices de Correção do Contrato Adimplemento"))
    assert adimplemento["Correção monetária"] == "IPCA"
    assert adimplemento["Juros remuneratórios"] == "1% ao mês"
    assert "Juros moratórios" not in adimplemento, "mora não incide enquanto adimplente"

    inadimplemento = _valores(_tabela(documento, "Índices de Correção do Contrato Inadimplemento"))
    assert inadimplemento["Correção monetária"] == "INPC"
    assert inadimplemento["Juros moratórios"] == "1% ao mês"
    assert inadimplemento["Comissão de permanência"] == "Vedada cumulação"


def testar_planilha_inicial_ganhou_comissao_de_permanencia(documento):
    valores = _valores(_tabela(documento, "Planilha Inicial"))
    assert valores["Correção monetária"] == "INPC"
    assert valores["Comissão de permanência"] == "Não aplicada"


def testar_redes_sociais_preenche_a_plataforma_certa(documento):
    """As plataformas do modelo ficam onde estão; perfil de plataforma fora da lista
    entra como linha nova, em vez de se perder."""
    tabela = _tabela(documento, "REDES SOCIAIS")
    linhas = {l.cells[0].text.strip(): (l.cells[1].text.strip(), l.cells[2].text.strip())
              for l in tabela.rows[1:]}
    assert linhas["Instagram"][0] == "@agroteste"
    assert linhas["Facebook"] == ("", ""), "plataforma sem perfil continua em branco"
    assert linhas["Kwai"][0] == "kwai.com/@joao"


def testar_redes_sociais_sem_perfil_nao_mexe_no_quadro():
    doc = Document(dossie_ppa._build_doc(dict(DADOS, redes_sociais=[])))
    tabela = dossie_ppa._tabela_por_titulo(doc, "REDES SOCIAIS")
    assert len(tabela.rows) == 8, "as 7 plataformas do modelo continuam lá, e nada além"
    assert all(not l.cells[1].text.strip() for l in tabela.rows[1:])


def testar_vp_aparece_como_vf(documento):
    """O modelo renomeou VP para VF; a chave da extração continua "vp"."""
    valores = _valores(documento.tables[1])
    assert valores["Total atingível mapeado — VF"] == "R$ 1.400.000,00"
    assert "Total atingível mapeado — VP" not in valores
    texto = _texto_inteiro(documento)
    assert "(VF)" in texto and "(VP)" not in texto


def testar_teses_sao_localizadas_por_titulo_e_nao_por_indice(documento):
    """REDES SOCIAIS e o quadro de lastro deslocaram todos os índices de tabela — o
    conteúdo das teses tem de continuar caindo no quadro certo."""
    assert "Penhora deferida sobre a sede." in _tabela(documento, "a. Ponderações").cell(0, 0).text
    caixas = [t for titulo, t in dossie_ppa._iter_headings_tables(documento)
              if titulo.strip().casefold().startswith("a. resumo da tese")]
    assert len(caixas) == 2, "2.2 (IDPJ) e 2.3 (fraude) têm cada uma a sua"
    assert "Confusão patrimonial" in caixas[0].cell(0, 0).text
    assert "Doação ao filho" in caixas[1].cell(0, 0).text
    empresa = _valores(_tabela(documento, "b. Empresa-Alvo"))
    assert empresa["Razão social da empresa-alvo"] == "Holding Teste Ltda"
    matriculas = _tabela(documento, "c. Matrículas Mapeadas")
    assert "30.174" in matriculas.rows[1].cells[0].text


def testar_consideracoes_gerais_caem_na_caixa_certa(documento):
    texto = _texto_inteiro(documento)
    assert "Execução com penhora deferida." in texto
    assert "Campo livre para registro" not in texto, "a orientação some quando há conteúdo"


def testar_visao_geral_dos_atingiveis_nao_engole_redes_sociais():
    """A clonagem por tese vai até o próximo capítulo — que agora é REDES SOCIAIS."""
    caminho = dossie_ppa.preencher_ativos_e_passivo_dossie(
        None,
        [{"tese": "Penhora Direta", "vm": "R$ 1,00"}],
        [{"tese": "Penhora Direta", "linhas": [["Matrícula nº 1", "", "", "", "", "", "", ""]]}],
        [], [], [], [],
    )
    doc = Document(caminho)
    assert dossie_ppa._tabela_por_titulo(doc, "REDES SOCIAIS") is not None
    assert "REDES SOCIAIS" in _texto_inteiro(doc)
